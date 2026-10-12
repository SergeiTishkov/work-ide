"""
A CV as a Word document: the content of a tailored CV -> .docx, and the text
of an existing .docx CV -> plain text for the model to read.

The layout copies the owner's own CV (2026-10-06): A4, 1.27 cm margins,
Arial 11 at 1.15 line spacing; the name, headline and stack line on the left
and the contacts in small print on the right; section titles as white bold
text on a teal bar; each job as a bold date line and a bold "Title, Company"
line, one line of context, bullets, and a "Tech stack:" line. LAYOUT holds the
numbers, so another person's look is a change in one place.

The content is a plain dict (see cv_tailor.CONTENT_SCHEMA):

    {"name": ..., "headline": ..., "tagline": ...,
     "contacts": [{"label": "Email", "value": "..."}],
     "summary": ["paragraph", ...],
     "skills": [{"category": "Back End", "items": ["C#", ...]}],
     "experience": [{"dates": ..., "title": ..., "company": ..., "context": ...,
                     "bullets": [{"text": ..., "priority": 1}], "tech_stack": [...]}],
     "sections": [{"heading": "Education", "lines": ["..."]}]}

A bullet may also be a plain string. `priority` is used by the length fitting
in cv_tailor and is not printed.

render_letter() writes a cover letter (cover_letter.LETTER_SCHEMA) under the
same header as the CV, so the two read as one set:

    {"name": ..., "headline": ..., "contacts": [...], "date": ...,
     "recipient": ["Hiring Team", "Contoso"], "greeting": "Dear Hiring Team,",
     "paragraphs": ["...", ...], "closing": "Kind regards,"}
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, List

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Cm, Pt, RGBColor

LAYOUT = {
    "page_width_cm": 21.0,
    "page_height_cm": 29.7,
    "margin_cm": 1.27,
    "font": "Arial",
    "font_size_pt": 11,
    "line_spacing": 1.15,
    "header_size_pt": 12,          # name, headline, stack line
    "contact_size_pt": 8,
    "contact_tab_cm": 3.0,         # where a contact's value starts
    "bar_fill": "31849B",          # section title bar
    "bar_text_size_pt": 12,
    "bar_text_indent_cm": 1.27,
    "bar_space_before_pt": 6,
    "bar_space_after_pt": 5,
    "job_head_size_pt": 14,        # date line and "Title, Company"
    "job_space_before_pt": 14,
    "block_gap_pt": 11,            # after the context line, after the bullets
    "summary_first_line_cm": 1.27,
    "bullet_indent_cm": 0.635,
    "tech_stack_label": "Tech stack:",
    "tech_stack_tab_cm": 2.5,
    "tech_stack_separator": " · ",
    "link_color": "0000FF",
    "section_titles": {"summary": "Summary", "skills": "Skills", "experience": "Work experience"},
}

URL_RE = re.compile(r"^(https?://\S+)$", re.IGNORECASE)
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^\+?[\d\s()\-]{7,}$")


# --- reading ----------------------------------------------------------------

def docx_text(path: Path) -> str:
    """The text of a .docx in reading order, tables included, one paragraph
    per line. Tabs survive as tabs: "Tech stack:\t.NET" stays one line."""
    document = Document(str(path))
    lines: List[str] = []
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            lines.append(_paragraph_text(child))
        elif child.tag == qn("w:tbl"):
            for paragraph in child.iter(qn("w:p")):
                lines.append(_paragraph_text(paragraph))
    text = "\n".join(line.rstrip() for line in lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"


def _paragraph_text(p) -> str:
    # Only what is inside runs: a <w:tab> in the paragraph properties is a tab
    # stop, not a tab character.
    parts = []
    for run in p.iter(qn("w:r")):
        for node in run:
            if node.tag == qn("w:t"):
                parts.append(node.text or "")
            elif node.tag == qn("w:tab"):
                parts.append("\t")
            elif node.tag in (qn("w:br"), qn("w:cr")) and node.get(qn("w:type")) in (None, "textWrapping"):
                parts.append("\n")
    return "".join(parts)


def read_cv_text(path: Path) -> str:
    """A CV file as text: .docx through docx_text, .md/.txt as they are."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return docx_text(path)
    if suffix in (".md", ".txt", ".markdown"):
        return path.read_text(encoding="utf-8")
    raise ValueError(f"cannot read a CV from {path.name}: .docx, .md or .txt expected")


# --- writing ----------------------------------------------------------------

def bullet_text(bullet) -> str:
    return bullet["text"] if isinstance(bullet, dict) else str(bullet)


def render(content: dict, path: Path, *, layout: dict = None, author: str = "") -> Path:
    """Writes the CV to `path` (.docx) and returns it."""
    lay = dict(LAYOUT, **(layout or {}))
    document = Document()
    _base_style(document, lay)
    _page(document.sections[0], lay)
    _header(document, content, lay)

    titles = lay["section_titles"]
    summary = [p for p in content.get("summary") or [] if p and p.strip()]
    if summary:
        _bar(document, titles["summary"], lay)
        for text in summary:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Cm(lay["summary_first_line_cm"])
            paragraph.paragraph_format.space_after = Pt(4)
            _run(paragraph, text.strip())

    skills = [s for s in content.get("skills") or [] if s.get("items")]
    if skills:
        _bar(document, titles["skills"], lay)
        for group in skills:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            _run(paragraph, f"{group.get('category', '').strip()}", bold=True)
            _run(paragraph, ": " + ", ".join(i.strip() for i in group["items"] if i.strip()))

    jobs = content.get("experience") or []
    if jobs:
        _bar(document, titles["experience"], lay)
        for job in jobs:
            _job(document, job, lay)

    for section in content.get("sections") or []:
        lines = [line for line in section.get("lines") or [] if line and line.strip()]
        if not section.get("heading") or not lines:
            continue
        _bar(document, section["heading"], lay)
        for line in lines:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            _run(paragraph, line.strip())

    properties = document.core_properties
    properties.author = author or content.get("name", "")
    properties.last_modified_by = properties.author
    properties.title = f"CV {content.get('name', '')}".strip()
    properties.comments = ""
    properties.subject = content.get("headline", "")
    properties.keywords = ""

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def render_letter(content: dict, path: Path, *, layout: dict = None, author: str = "") -> Path:
    """Writes the cover letter to `path` (.docx) and returns it."""
    lay = dict(LAYOUT, **(layout or {}))
    document = Document()
    _base_style(document, lay)
    _page(document.sections[0], lay)
    _header(document, content, lay)

    def block(lines, space_before=0.0, space_after=0.0):
        lines = [line.strip() for line in lines if line and line.strip()]
        for index, line in enumerate(lines):
            paragraph = document.add_paragraph()
            if index == 0:
                paragraph.paragraph_format.space_before = Pt(space_before)
            if index == len(lines) - 1:
                paragraph.paragraph_format.space_after = Pt(space_after)
            _run(paragraph, line)

    gap = lay["block_gap_pt"]
    block([content.get("date", "")], space_before=2 * gap, space_after=gap)
    block(content.get("recipient") or [], space_after=gap)
    block([content.get("greeting", "")], space_after=gap)
    for text in content.get("paragraphs") or []:
        block([text], space_after=gap)
    block([content.get("closing", "")])
    block([signature(content.get("name", ""))], space_before=gap)

    properties = document.core_properties
    properties.author = author or content.get("name", "")
    properties.last_modified_by = properties.author
    properties.title = f"Cover letter {content.get('name', '')}".strip()
    properties.comments = ""
    properties.subject = content.get("headline", "")
    properties.keywords = ""

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def signature(name: str) -> str:
    """The name under a letter: as a person signs it. A CV header often has
    it in capitals ("ALEX EXAMPLE"); those become "Alex Example"."""
    name = (name or "").strip()
    return name.title() if name.isupper() else name


def _header(document, content: dict, lay: dict) -> None:
    """The header: two columns, as in the original. Left: name, headline,
    stack line. Right: contacts. The body after it is one column."""
    _set_columns(document.sections[0], 2)
    for line in (content.get("name", ""), content.get("headline", ""), content.get("tagline", "")):
        if line:
            paragraph = document.add_paragraph()
            _run(paragraph, line, bold=True, size=lay["header_size_pt"])
    contacts = content.get("contacts") or []
    for index, contact in enumerate(contacts):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(lay["contact_tab_cm"]))
        if index == 0:
            paragraph.add_run().add_break(WD_BREAK.COLUMN)
        label = (contact.get("label") or "").strip()
        if label:
            _run(paragraph, f"{label.rstrip(':')}:\t", size=lay["contact_size_pt"])
        _contact_value(paragraph, (contact.get("value") or "").strip(), lay)

    body = document.add_section(WD_SECTION.CONTINUOUS)
    _page(body, lay)
    _set_columns(body, 1)


def _job(document, job: dict, lay: dict) -> None:
    dates = (job.get("dates") or "").strip()
    head = ", ".join(x for x in ((job.get("title") or "").strip(), (job.get("company") or "").strip()) if x)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(lay["job_space_before_pt"])
    paragraph.paragraph_format.keep_with_next = True
    _run(paragraph, f"{dates}:" if dates else head, bold=True, size=lay["job_head_size_pt"])
    if dates and head:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        _run(paragraph, head, bold=True, size=lay["job_head_size_pt"])
    context = (job.get("context") or "").strip()
    bullets = [bullet_text(b).strip() for b in job.get("bullets") or [] if bullet_text(b).strip()]
    if context:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(lay["block_gap_pt"])
        paragraph.paragraph_format.keep_with_next = bool(bullets)
        _run(paragraph, context)
    for index, text in enumerate(bullets):
        paragraph = document.add_paragraph(style="List Bullet")
        fmt = paragraph.paragraph_format
        fmt.left_indent = Cm(lay["bullet_indent_cm"])
        fmt.first_line_indent = Cm(-lay["bullet_indent_cm"])
        fmt.line_spacing = 1.0
        if index == len(bullets) - 1:
            fmt.space_after = Pt(lay["block_gap_pt"])
        _run(paragraph, text)
    stack = [t.strip() for t in job.get("tech_stack") or [] if t and t.strip()]
    if stack:
        paragraph = document.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.tab_stops.add_tab_stop(Cm(lay["tech_stack_tab_cm"]))
        fmt.left_indent = Cm(lay["tech_stack_tab_cm"])
        fmt.first_line_indent = Cm(-lay["tech_stack_tab_cm"])
        fmt.line_spacing = 1.0
        _run(paragraph, f"{lay['tech_stack_label']}\t" + lay["tech_stack_separator"].join(stack))


def _base_style(document, lay: dict) -> None:
    normal = document.styles["Normal"]
    normal.font.name = lay["font"]
    normal.font.size = Pt(lay["font_size_pt"])
    fonts = normal.element.rPr.find(qn("w:rFonts"))
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), lay["font"])
    fmt = normal.paragraph_format
    fmt.line_spacing = lay["line_spacing"]
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    # The default template gives List Bullet "contextual spacing" and its own
    # indents; the indents are set per paragraph in _job.
    bullet = document.styles["List Bullet"]
    bullet.font.name = lay["font"]


def _page(section, lay: dict) -> None:
    section.page_width = Cm(lay["page_width_cm"])
    section.page_height = Cm(lay["page_height_cm"])
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(lay["margin_cm"]))
    section.header_distance = Cm(0)


def _set_columns(section, number: int) -> None:
    sect = section._sectPr
    cols = sect.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect.append(cols)
    cols.set(qn("w:num"), str(number))
    cols.set(qn("w:space"), "720")


def _run(paragraph, text: str, *, bold: bool = False, size: float = None, color: str = None):
    run = paragraph.add_run(text)
    run.bold = bold or None
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _contact_value(paragraph, value: str, lay: dict) -> None:
    target = None
    if URL_RE.match(value):
        target = value
    elif EMAIL_RE.match(value):
        target = f"mailto:{value}"
    elif PHONE_RE.match(value):
        target = "tel:" + re.sub(r"[^\d+]", "", value)
    if not target:
        _run(paragraph, value, size=lay["contact_size_pt"])
        return
    relation = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), lay["link_color"])
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), str(int(lay["contact_size_pt"] * 2)))
    properties.extend([color, underline, size])
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = value
    text.set(qn("xml:space"), "preserve")
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def _bar(document, title: str, lay: dict) -> None:
    """A section title: white bold text on a full-width filled bar (a one-cell
    table, as in the original), with a little air above it."""
    spacer = document.add_paragraph()
    spacer.paragraph_format.line_spacing = Pt(lay["bar_space_before_pt"])
    spacer.paragraph_format.keep_with_next = True

    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    width = Cm(lay["page_width_cm"] - 2 * lay["margin_cm"])
    cell = table.rows[0].cells[0]
    cell.width = width
    tcPr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), lay["bar_fill"])
    tcPr.append(shading)
    margins = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "80")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tcPr.append(margins)
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tblPr.append(borders)
    # The heading row stays with the paragraph after it.
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.left_indent = Cm(lay["bar_text_indent_cm"])
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.line_spacing = 1.0
    _run(paragraph, title, bold=True, size=lay["bar_text_size_pt"], color="FFFFFF")
    spacer = document.add_paragraph()
    spacer.paragraph_format.line_spacing = Pt(lay["bar_space_after_pt"])
    spacer.paragraph_format.keep_with_next = True


def plain_text(content: dict) -> str:
    """The CV as plain text, in the order it is printed: for the fact check
    and for the length estimate."""
    lines: List[str] = [content.get("name", ""), content.get("headline", ""), content.get("tagline", "")]
    lines += [f"{c.get('label', '')}: {c.get('value', '')}" for c in content.get("contacts") or []]
    lines += list(content.get("summary") or [])
    lines += [f"{g.get('category', '')}: {', '.join(g.get('items') or [])}" for g in content.get("skills") or []]
    for job in content.get("experience") or []:
        lines += [f"{job.get('dates', '')}:", f"{job.get('title', '')}, {job.get('company', '')}",
                  job.get("context", "")]
        lines += [bullet_text(b) for b in job.get("bullets") or []]
        lines.append("Tech stack: " + " · ".join(job.get("tech_stack") or []))
    for section in content.get("sections") or []:
        lines.append(section.get("heading", ""))
        lines += list(section.get("lines") or [])
    return "\n".join(line for line in lines if line)


def iter_terms(content: dict) -> Iterable[str]:
    """Every skill and tech-stack entry: the claims the fact check matches
    against the source CVs one by one."""
    for group in content.get("skills") or []:
        for item in group.get("items") or []:
            yield item
    for job in content.get("experience") or []:
        for item in job.get("tech_stack") or []:
            yield item
